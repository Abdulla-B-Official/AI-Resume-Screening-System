#Handles file ingestion and text extraction for PDF, DOCX, and TXT files.
import os
import fitz  # PyMuPDF
from docx import Document

def extract_text_from_pdf(pdf_path):
    doc = fitz.open(pdf_path)
    pages = [page.get_text() for page in doc]
    doc.close()
    return "\n".join(pages)

def extract_text_from_docx(docx_path):
    document = Document(docx_path)
    parts = [para.text for para in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)

def extract_text_from_txt(txt_path):
    with open(txt_path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()

def extract_text(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext == ".docx":
        return extract_text_from_docx(file_path)
    elif ext == ".txt":
        return extract_text_from_txt(file_path)
    else:
        return ""